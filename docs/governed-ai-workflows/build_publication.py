from pathlib import Path
import os, re, textwrap
os.environ.setdefault('MPLCONFIGDIR', '/tmp/agent-paper-mpl')
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
FIG=ROOT/'figures'; FIG.mkdir(exist_ok=True)
OUT=ROOT/'Agent_Orchestration_Architecture_White_Paper.docx'
BLUE='#23485c'; PALE='#edf3f6'; LINE='#657581'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9})

def canvas(h):
    f,a=plt.subplots(figsize=(6.8,h)); f.subplots_adjust(0,0,1,1)
    a.set_xlim(0,6.8); a.set_ylim(0,h); a.axis('off'); return f,a
def box(a,x,y,w,h,t,fill=PALE):
    a.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.025,rounding_size=0.055',lw=.85,edgecolor=LINE,facecolor=fill))
    a.text(x+w/2,y+h/2,t,ha='center',va='center',fontsize=9,color='#132a36',linespacing=1.3)
def arrow(a,start,end,label=None,rad=0,dashed=False):
    a.add_patch(FancyArrowPatch(start,end,arrowstyle='-|>',mutation_scale=9,lw=.95,color=LINE,connectionstyle=f'arc3,rad={rad}',linestyle='--' if dashed else '-'))
    if label:
        x=(start[0]+end[0])/2; y=(start[1]+end[1])/2
        a.text(x,y+.07,label,ha='center',va='bottom',fontsize=8,color='#263843',bbox={'facecolor':'white','edgecolor':'none','pad':1})
def save(f,name):
    f.savefig(FIG/f'{name}.png',dpi=240,facecolor='white')
    f.savefig(FIG/f'{name}.svg',facecolor='white'); plt.close(f)

f,a=canvas(4.8)
box(a,.2,4.13,6.4,.48,'Control plane  |  approved policies, capabilities and release versions')
box(a,2.3,3.1,2.2,.6,'Durable coordinator\nTask state and budgets')
box(a,.2,3.1,1.6,.6,'Run ledger\nand artifacts')
box(a,5,3.1,1.6,.6,'Operator\nReview and stop')
box(a,2.3,2.12,2.2,.6,'Agent workers\nBounded reasoning')
box(a,.2,2.12,1.6,.6,'Context and\nreviewed memory')
box(a,2.3,1.14,2.2,.6,'Policy and tool gateway\nAuthorize each action')
box(a,2.3,.16,2.2,.6,'Execution adapters\nTools and sandboxes')
box(a,5,1.14,1.6,.6,'Independent\nverification')
arrow(a,(3.4,4.12),(3.4,3.72)); arrow(a,(1.8,3.4),(2.28,3.4)); arrow(a,(5,3.4),(4.52,3.4))
arrow(a,(3.4,3.1),(3.4,2.74)); arrow(a,(1.8,2.42),(2.28,2.42))
arrow(a,(3.4,2.12),(3.4,1.76)); arrow(a,(3.4,1.14),(3.4,.78))
arrow(a,(4.52,.46),(5.8,1.12),rad=.2); arrow(a,(5.8,1.77),(4.53,3.17),rad=.25)
a.text(6.48,2.22,'Evidence\nand feedback',ha='right',fontsize=8,color=LINE)
save(f,'architecture')

f,a=canvas(2.85)
for x,t in [(.15,'Verified observations\nSource references'),(2.45,'Candidate memory\nScope and validity'),(4.75,'Review\nResolve conflicts')]:box(a,x,1.85,1.9,.62,t)
for x,t in [(.15,'Bounded context\nFor this task'),(2.45,'Authorized recall\nFreshness checks'),(4.75,'Approved lessons\nVersioned records')]:box(a,x,.43,1.9,.62,t)
arrow(a,(2.08,2.16),(2.42,2.16));arrow(a,(4.38,2.16),(4.72,2.16));arrow(a,(5.7,1.82),(5.7,1.08))
arrow(a,(4.72,.74),(4.38,.74));arrow(a,(2.42,.74),(2.08,.74))
a.text(3.4,1.38,'Correction, expiry and revocation invalidate derived material',ha='center',fontsize=8.5,color=LINE)
save(f,'memory')

f,a=canvas(3.55)
xs=[.65,2.35,4.15,6.05]
for x,t in zip(xs,['Worker','Gateway','Adapter','Destination']):
    box(a,x-.55,3.04,1.1,.36,t)
    a.plot([x,x],[.15,3.02],color='#bec8ce',lw=.8,linestyle='--')
def msg(i,j,y,t):
    arrow(a,(xs[i],y),(xs[j],y));a.text((xs[i]+xs[j])/2,y+.08,t,ha='center',fontsize=8)
msg(0,1,2.7,'Propose exact action')
a.text(2.35,2.28,'Policy + approval\nLease + budget',ha='center',fontsize=8.5,color=BLUE)
msg(1,2,1.95,'Grant + operation ID')
msg(2,3,1.48,'Apply with preconditions')
msg(3,2,1.02,'Result or unknown')
msg(2,1,.60,'Record effect receipt')
msg(1,0,.21,'Return status and evidence')
save(f,'action')

f,a=canvas(2.4)
a.set_ylim(0,3.25)
nodes={'admitted':(.15,2.52,'Admitted'),'ready':(2.5,2.52,'Ready'),'running':(4.85,2.52,'Running'),
       'waiting':(.15,1.48,'Waiting'),'verify':(2.5,1.48,'Verifying'),'reconcile':(4.85,1.48,'Reconciling'),
       'canceled':(.15,.35,'Canceled'),'done':(2.5,.35,'Completed'),'unresolved':(4.85,.35,'Unresolved')}
for _,(x,y,t) in nodes.items():box(a,x,y,1.8,.48,t)
arrow(a,(1.97,2.76),(2.47,2.76));arrow(a,(4.32,2.76),(4.82,2.76));arrow(a,(4.9,2.51),(4.18,1.98))
arrow(a,(5.75,2.51),(5.75,1.99));arrow(a,(5.75,1.46),(5.75,.85))
arrow(a,(3.4,1.46),(3.4,.85));arrow(a,(3.4,1.98),(3.4,2.5))
arrow(a,(2.48,1.72),(1.98,1.72));arrow(a,(1.05,1.98),(2.47,2.61),rad=.1)
arrow(a,(1.05,1.46),(1.05,.85));arrow(a,(4.83,1.72),(4.32,1.72))
a.text(3.4,.05,'Failed is a terminal alternative; unresolved work requires reconciliation.',ha='center',fontsize=8,color=LINE)
save(f,'lifecycle')

f,a=canvas(3.0)
for x,t in [(.15,'Production outcomes\nPermitted evidence'),(2.45,'Candidate change\nExplicit hypothesis'),(4.75,'Held-out evaluation\nChallenge the claim')]:box(a,x,2.05,1.9,.65,t)
for x,t in [(.15,'Baseline release\nRollback available'),(2.45,'Small cohort\nObserve outcomes'),(4.75,'Independent approval\nPinned release bundle')]:box(a,x,.5,1.9,.65,t)
arrow(a,(2.08,2.37),(2.42,2.37));arrow(a,(4.38,2.37),(4.72,2.37));arrow(a,(5.7,2.02),(5.7,1.18))
arrow(a,(4.72,.82),(4.38,.82));arrow(a,(2.42,.82),(2.08,.82));arrow(a,(1.1,1.18),(1.1,2.02))
a.text(3.4,1.52,'Permission and evaluator changes require separate controlled review',ha='center',fontsize=8.4,color=LINE)
save(f,'improvement')

f,a=canvas(5.1)
box(a,.2,4.35,2.8,.55,'Developer or agent\nDelivery intent and source')
box(a,3.8,4.35,2.8,.55,'Approved registry\nCapabilities and requirements')
box(a,1.9,3.35,3,.60,'Plan compiler and coordinator\nValidate and schedule work')
box(a,.2,2.25,2.0,.65,'Isolated build\nImmutable artifact')
box(a,2.4,2.25,2.0,.65,'Protected verifier\nArtifact evidence')
box(a,4.6,2.25,2.0,.65,'Release authorizer\nScoped grant')
box(a,4.6,1.1,2.0,.65,'Target controller\nApply or reconcile')
box(a,2.4,1.1,2.0,.65,'Health observer\nPass fail or wait')
box(a,.2,1.1,2.0,.65,'Recovery controller\nBounded response')
arrow(a,(1.6,4.32),(2.6,3.98));arrow(a,(5.2,4.32),(4.2,3.98))
arrow(a,(2.5,3.33),(1.2,2.93));arrow(a,(3.4,3.33),(3.4,2.93))
arrow(a,(2.22,2.58),(2.37,2.58));arrow(a,(4.42,2.58),(4.57,2.58))
arrow(a,(5.6,2.23),(5.6,1.78));arrow(a,(4.57,1.42),(4.42,1.42))
arrow(a,(2.37,1.42),(2.22,1.42))
box(a,.2,.16,6.4,.45,'Durable ledger and evidence store  |  every transition has a record')
a.text(3.4,.81,'AI proposes work; trusted services control consequential effects.',ha='center',fontsize=8.5)
save(f,'delivery')

f,a=canvas(3.6)
for x,t in [(.2,'Candidate'),(2.5,'Evidence ready'),(4.8,'Authorized')]:box(a,x,2.7,1.8,.55,t)
for x,t in [(.2,'Recovered'),(2.5,'Observing'),(4.8,'Applying')]:box(a,x,1.55,1.8,.55,t)
box(a,.2,.3,1.8,.55,'Failed or blocked');box(a,2.5,.3,1.8,.55,'Accepted');box(a,4.8,.3,1.8,.55,'Unknown effect')
arrow(a,(2.02,2.98),(2.47,2.98));arrow(a,(4.32,2.98),(4.77,2.98))
arrow(a,(5.7,2.68),(5.7,2.13));arrow(a,(4.77,1.83),(4.32,1.83))
arrow(a,(3.4,1.53),(3.4,.88),'pass');arrow(a,(2.47,1.83),(2.02,1.83),'fail')
arrow(a,(5.7,1.53),(5.7,.88));arrow(a,(4.78,.58),(4.3,1.56),rad=-.2)
arrow(a,(1.1,1.53),(1.1,.88))
a.text(3.4,3.43,'Changed evidence or target state invalidates authorization.',ha='center',fontsize=8.5)
save(f,'release')

f,a=canvas(3.5)
for x,t in [(.2,'1 Observe\nInventory and baseline'),(2.5,'2 Wrap\nCommon contracts'),(4.8,'3 Replace execution\nTrusted capabilities')]:box(a,x,2.35,1.8,.7,t)
for x,t in [(.2,'6 Retire\nAll obligations covered'),(2.5,'5 Disconnect\nProve independence'),(4.8,'4 Replace release\nOne target owner')]:box(a,x,.7,1.8,.7,t)
arrow(a,(2.02,2.7),(2.47,2.7));arrow(a,(4.32,2.7),(4.77,2.7));arrow(a,(5.7,2.33),(5.7,1.43))
arrow(a,(4.77,1.05),(4.32,1.05));arrow(a,(2.47,1.05),(2.02,1.05))
a.text(3.4,1.83,'Advance by evidence for each application cohort.',ha='center',fontsize=9)
a.text(3.4,.24,'Keep the current path when a required guarantee or economic case is missing.',ha='center',fontsize=8.3)
save(f,'migration')

f,a=canvas(3.8)
for x,t in [(.2,'Verified outcomes\nIncluding failures'),(2.5,'Cases and facts\nScoped and current'),(4.8,'Obligations and tests\nIndependent owners')]:box(a,x,2.55,1.8,.7,t)
for x,t in [(.2,'Replace or retain\nPortable contracts'),(2.5,'Fair comparison\nSame authorized inputs'),(4.8,'Candidate capability\nNative or custom')]:box(a,x,.8,1.8,.7,t)
arrow(a,(2.02,2.9),(2.47,2.9));arrow(a,(4.32,2.9),(4.77,2.9));arrow(a,(5.7,2.53),(5.7,1.53))
arrow(a,(4.77,1.15),(4.32,1.15));arrow(a,(2.47,1.15),(2.02,1.15));arrow(a,(1.1,1.53),(1.1,2.53))
a.text(3.4,2,'Proposed lessons and requirement changes require review.',ha='center',fontsize=8.5)
a.text(3.4,.28,'Retain custom implementation only while measured outcomes justify it.',ha='center',fontsize=8.5)
save(f,'value')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.9);sec.left_margin=Inches(.8);sec.right_margin=Inches(.8)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.3)
styles=doc.styles
for name in ['Normal','Body Text']:
    s=styles[name];s.font.name='Calibri';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=1.05
for name,size in [('Title',28),('Subtitle',15),('Heading 1',19),('Heading 2',14),('Heading 3',12)]:
    s=styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(10 if name!='Title' else 14);s.paragraph_format.space_after=Pt(8)
    s.paragraph_format.keep_with_next=True
styles['Caption'].font.size=Pt(9);styles['Caption'].font.italic=False;styles['Caption'].font.color.rgb=RGBColor(55,55,55)
styles['Caption'].paragraph_format.space_after=Pt(9)
for s in ['Normal','Heading 1','Heading 2','Heading 3']:
    styles[s].paragraph_format.widow_control=True
for s in styles:
    for border in list(s.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Agent orchestration  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='Governed AI Workflows and the Future of Software Delivery'
doc.core_properties.subject='Reference architecture and implementation blueprint'
doc.core_properties.author='';doc.core_properties.keywords='AI agents, orchestration, durable execution, memory, evaluation'

def inline(p,t):
    parts=re.split(r'(\*\*.*?\*\*|`[^`]+`)',t)
    for part in parts:
        if part.startswith('**'):
            r=p.add_run(part[2:-2]);r.bold=True
        elif part.startswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(9)
        else:p.add_run(part)

def table(lines):
    rows=[[c.strip() for c in l.strip().strip('|').split('|')] for l in lines]
    rows=[r for r in rows if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
    n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths={2:[1.85,5.05],3:[1.55,2.8,2.55]}.get(n,[6.9/n]*n)
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        trPr=t.rows[ri]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for ci,(cell,content) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'23485C' if ri==0 else ('F1F5F7' if ri%2==0 else 'FFFFFF'));tcPr.append(sh)
            mar=OxmlElement('w:tcMar')
            for tag in ['top','left','bottom','right']:
                e=OxmlElement('w:'+tag);e.set(qn('w:w'),'60' if tag in ['top','bottom'] else '80');e.set(qn('w:type'),'dxa');mar.append(e)
            tcPr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for tag in ['top','left','bottom','right']:
                b=OxmlElement('w:'+tag);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            tcPr.append(borders)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1.5);p.paragraph_format.space_before=Pt(1.5);p.paragraph_format.line_spacing=1.0
            inline(p,content)
            for r in p.runs:
                r.font.size=Pt(9.5)
                if ri==0:r.bold=True;r.font.color.rgb=RGBColor(255,255,255)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=Pt(1)
    p.add_run().font.size=Pt(1)

lines=(ROOT/'whitepaper.md').read_text().splitlines();i=0;first=True;newpage=False
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line=='<!-- page -->':newpage=True;i+=1;continue
    if line.startswith('```'):
        code=[];i+=1
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        for l in code:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            p.paragraph_format.keep_with_next=True
            r=p.add_run(l);r.font.name='Consolas';r.font.size=Pt(8.5)
        if code:p.paragraph_format.keep_with_next=False
        doc.add_paragraph().paragraph_format.space_after=Pt(1)
        i+=1;continue
    if line.startswith('|'):
        block=[]
        while i<len(lines) and lines[i].strip().startswith('|'):block.append(lines[i]);i+=1
        table(block);continue
    if line.startswith('!['):
        m=re.match(r'!\[(.*?)\]\((.*?)\)',line)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        shape=p.add_run().add_picture(str(ROOT/m.group(2)),width=Inches(6.8))
        shape._inline.docPr.set('descr',m.group(1))
        p=doc.add_paragraph(m.group(1),'Caption');i+=1;continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:], 'Title' if first else 'Heading 1');first=False
        if newpage:p.paragraph_format.page_break_before=True;newpage=False
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Subtitle')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif re.match(r'^\d+\. ',line):
        p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);inline(p,line)
    else:
        p=doc.add_paragraph();inline(p,line)
        if line.startswith('['):
            p.paragraph_format.space_after=Pt(10)
            for r in p.runs:r.font.size=Pt(9.5)
    i+=1
doc.save(OUT)
print(OUT)

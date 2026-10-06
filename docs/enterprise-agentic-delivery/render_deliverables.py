from pathlib import Path
import re, math, textwrap, json, zipfile, io, sys
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from reportlab.platypus.tableofcontents import TableOfContents
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from xml.sax.saxutils import escape
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement

ROOT=Path(__file__).parent
ASSETS=ROOT/'diagrams'; ASSETS.mkdir(exist_ok=True)
NAVY='#132B45'; BLUE='#21618C'; TEAL='#087F8C'; LIGHT='#EDF4F8'; GOLD='#E5AB46'; GRAY='#536575'

def diagram(name,title,subtitle,nodes,edges,groups=(),size=(12,10)):
    fig,ax=plt.subplots(figsize=size)
    fig.patch.set_facecolor('white'); ax.set_xlim(0,12); ax.set_ylim(0,10); ax.axis('off')
    ax.text(.3,9.7,title,fontsize=18,fontweight='bold',color=NAVY,va='top')
    ax.text(.3,9.16,subtitle,fontsize=9.5,color=GRAY,va='top')
    for x,y,w,h,label in groups:
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.06',fc='#F7F9FB',ec='#A8BBCB',ls='--',lw=1.2,zorder=0))
        ax.text(x+.14,y+h-.16,label,fontsize=9.5,color=GRAY,va='top')
    for a,b,label,kind in edges:
        na,nb=nodes[a],nodes[b]
        x1,y1,w1,h1=na[:4]; x2,y2,w2,h2=nb[:4]
        c1=(x1+w1/2,y1+h1/2); c2=(x2+w2/2,y2+h2/2)
        dx,dy=c2[0]-c1[0],c2[1]-c1[1]
        def bound(c,dx,dy,w,h):
            t=min((w/2)/abs(dx) if dx else 1e5,(h/2)/abs(dy) if dy else 1e5)
            return c[0]+dx*t,c[1]+dy*t
        p1=bound(c1,dx,dy,w1,h1); p2=bound(c2,-dx,-dy,w2,h2)
        if kind=='feedback':
            p1=(x1+w1,y1+h1/2); p2=(x2+w2,y2+h2/2)
            ax.plot([p1[0],8.15,8.15,p2[0]],[p1[1],p1[1],p2[1],p2[1]],color=GRAY,lw=1.3,zorder=1)
            ax.add_patch(FancyArrowPatch((8.15,p2[1]),p2,arrowstyle='-|>',mutation_scale=13,lw=1.3,color=GRAY,zorder=1))
            ax.text(8.15,4.45,label,fontsize=7,color=GRAY,ha='center',va='center',bbox=dict(fc='white',ec='none',pad=2),zorder=4)
            continue
        ax.add_patch(FancyArrowPatch(p1,p2,arrowstyle='<->' if kind=='bidirectional' else '-|>',mutation_scale=13,lw=1.3,color=GRAY,ls='--' if kind=='optional' else '-',zorder=1,connectionstyle='arc3,rad=0.04'))
        if label:
            mx,my=(p1[0]+p2[0])/2,(p1[1]+p2[1])/2
            ax.text(mx,my,label,fontsize=8,color=GRAY,ha='center',va='center',bbox=dict(fc='white',ec='none',pad=2),zorder=4)
    for k,n in nodes.items():
        x,y,w,h,label,typ=n
        face=TEAL if typ=='core' else BLUE if typ=='module' else LIGHT
        ink='white' if typ in ('core','module') else NAVY
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.03',fc=face,ec=BLUE if typ=='external' else face,lw=1.1,zorder=2))
        label='\n'.join(textwrap.fill(line,width=int(w*10),break_long_words=False) for line in label.split('\n'))
        ax.text(x+w/2,y+h/2,label,fontsize=10.0,ha='center',va='center',color=ink,linespacing=1.3,zorder=3)
    ax.text(.3,.15,'Solid arrows: interaction/dependency   |   Dashed: optional scheduling or logical boundary',fontsize=8,color=GRAY)
    png=io.BytesIO()
    fig.savefig(png,format='png',dpi=190,bbox_inches='tight',facecolor='white')
    payload=png.getvalue()
    assert len(payload)>1000
    (ASSETS/f'{name}.png').write_bytes(payload)
    fig.savefig(ASSETS/f'{name}.svg',bbox_inches='tight',facecolor='white')
    plt.close(fig)

diagram('C1_Context','C1 | System context','Who uses the system, and which external systems it relies on.',{
 'owner':(4.4,7.6,3.2,1.0,'Engineer / service owner\nOperator','external'),
 'system':(3.9,4.4,4.2,1.45,'Enterprise Agentic Delivery\nOutcome, evidence and action contracts','core'),
 'harness':(.5,1.7,3.3,1.35,'Approved agent harnesses\nReasoning and artifact creation','external'),
 'identity':(8.2,1.7,3.3,1.35,'Enterprise identity and policy\nAuthority and scoped access','external'),
 'delivery':(4.2,.7,3.6,1.25,'Repository, build, release\nand operations systems','external')},
 [('owner','system','Define / inspect','bidirectional'),('system','harness','Tasks / results','bidirectional'),('system','identity','Policy decisions','bidirectional'),('system','delivery','Actions / evidence','bidirectional')])

diagram('C2_Containers','C2 | Containers and trust boundaries','Target implementation; reuse existing services for these roles wherever possible.',{
 'client':(4.3,7.6,3.4,.8,'Existing IDE / portal / events','external'),
 'api':(4.2,5.65,3.6,1.1,'Outcome Service / API\nCross-system records and decisions','core'),
 'worker':(.5,5.3,2.9,1.0,'Worker Gateway\nHarness adapters','module'),
 'runtime':(.5,2.7,2.9,1.3,'Approved harness\nIsolated workspace\nNo production credentials','external'),
 'record':(8.6,6.6,2.9,.9,'Outcome database\nState / effects / outbox','external'),
 'evidence':(8.6,4.9,2.9,.9,'Evidence object store\nProtected artifact records','external'),
 'action':(4.2,2.85,3.6,1.1,'Trusted Action Gateway\nAuthorize / execute / reconcile','core'),
 'policy':(8.6,2.85,2.9,1.1,'IAM / policy authority\nCurrent scoped decisions','external'),
 'target':(4.2,.6,3.6,1.1,'Repository / release executors\nEvidence and service observations','external'),
 'durable':(8.6,.6,2.9,1.1,'Optional Workflow Runner\nDurable cross-system scheduling','external')},
 [('client','api','',''),('api','worker','',''),('worker','runtime','',''),('runtime','api','Artifacts',''),('api','record','',''),('api','evidence','',''),('api','action','Bound action',''),('action','policy','',''),('action','target','',''),('target','api','Evidence /\nobservations','feedback'),('api','durable','','optional')],
 [(.28,2.38,3.35,4.4,'Untrusted reasoning'),(3.97,.34,4.06,4.3,'Trusted effects')])

diagram('C3_Components','C3 | Inside the Outcome Service','Deterministic authority and transitions surround flexible planning.',{
 'intake':(4.2,7.45,3.6,.95,'Intake and Contract Validator','module'),
 'transition':(4.2,5.55,3.6,1.15,'Transition Engine\nState and readiness rules','core'),
 'planner':(.5,5.25,3.0,1.1,'Planning and Harness Adapter\nProposals / bounded tasks','module'),
 'verifier':(.5,2.95,3.0,1.1,'Evidence Validator\nBinding / provenance / freshness','module'),
 'authority':(8.5,5.25,3.0,1.1,'Policy and Approval Connector\nCurrent authority decisions','module'),
 'effects':(8.5,2.95,3.0,1.1,'Action Dispatcher and Reconciler\nExternal operation identity','module'),
 'records':(4.2,2.85,3.6,1.0,'Record Repository and Outbox\nVersioned durable state','module'),
 'view':(4.2,.75,3.6,.95,'Status Projection and Telemetry','module')},
 [('intake','transition','',''),('transition','planner','',''),('planner','verifier','',''),('verifier','transition','',''),('transition','authority','',''),('authority','effects','',''),('effects','transition','',''),('transition','records','',''),('records','view','','')],
 [(.25,.45,11.5,8.47,'Container: Outcome Service')])

diagram('C4_Code','C4 | Code view: bounded external effects','Planned types/interfaces inside the C3 Action Dispatcher and Reconciler.',{
 'dispatch':(4.15,4.5,3.7,1.7,'ActionDispatcher\nsubmit(request): OperationRef\nreconcile(id): ActionResult','core'),
 'request':(4.1,7.4,3.8,1.2,'ActionRequest\noperationId / artifactDigest\nexpectedVersion / decisionId','external'),
 'policy':(.4,2.0,3.4,1.65,'PolicyAuthority\nauthorize(request, now)\nreturns Decision','external'),
 'executor':(4.3,.65,3.4,1.65,'ActionExecutor\napply(request, fence)\ninspect(operationId)','external'),
 'repo':(8.2,2.0,3.4,1.65,'EffectRepository\nreserve(requestHash)\nrecord(result, expectedVersion)','external')},
 [('dispatch','request','uses',''),('dispatch','policy','depends on',''),('dispatch','executor','depends on',''),('dispatch','repo','depends on','')])

FONT='/usr/share/fonts/truetype/dejavu/'
for name,file in [('Body','DejaVuSans.ttf'),('Body-Bold','DejaVuSans-Bold.ttf'),('Body-Italic','DejaVuSans.ttf'),('Mono','DejaVuSansMono.ttf')]:
    pdfmetrics.registerFont(TTFont(name,FONT+file))
pdfmetrics.registerFontFamily('Body',normal='Body',bold='Body-Bold',italic='Body-Italic',boldItalic='Body-Bold')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='Text',fontName='Body',fontSize=9.2,leading=14.1,spaceAfter=8,textColor=colors.HexColor(NAVY)))
styles.add(ParagraphStyle(name='Chapter',fontName='Body-Bold',fontSize=18,leading=23,spaceAfter=15,keepWithNext=True,textColor=colors.HexColor(NAVY)))
styles.add(ParagraphStyle(name='Subhead',fontName='Body-Bold',fontSize=12,leading=17,spaceBefore=13,spaceAfter=8,keepWithNext=True,textColor=colors.HexColor(TEAL)))
styles.add(ParagraphStyle(name='Cell',fontName='Body',fontSize=7.5,leading=11,textColor=colors.HexColor(NAVY)))
styles.add(ParagraphStyle(name='ContractCode',fontName='Mono',fontSize=7,leading=10,spaceAfter=1,textColor=colors.HexColor(NAVY)))
styles.add(ParagraphStyle(name='TOCEntry',fontName='Body',fontSize=9,leading=15,leftIndent=0,firstLineIndent=0,spaceBefore=4,textColor=colors.HexColor(NAVY)))

def inline(t):
    t=escape(t)
    t=re.sub(r'\*\*(.*?)\*\*',r'<b>\1</b>',t)
    t=re.sub(r'`(.*?)`',r'<font name="Mono">\1</font>',t)
    t=re.sub(r'(https://[^\s<]+)',lambda m:'<link href="'+m[0]+'" color="#21618C">'+m[0]+'</link>',t)
    return t

class PaperDoc(SimpleDocTemplate):
    def afterFlowable(self,f):
        if isinstance(f,Paragraph) and f.style.name=='Chapter':
            title=f.getPlainText(); key='h'+str(self.seq.nextf('heading'))
            if title=='Contents': return
            self.canv.bookmarkPage(key); self.canv.addOutlineEntry(title,key,0)
            self.notify('TOCEntry',(0,title,self.page,key))

def page(c,d):
    w,h=A4
    c.setStrokeColor(colors.HexColor('#CCD8E2')); c.line(48,h-39,w-48,h-39)
    c.setFont('Body',7); c.setFillColor(colors.HexColor(GRAY))
    c.drawString(48,h-29,'ENTERPRISE AGENTIC DELIVERY  /  ARCHITECTURE PROPOSAL')
    c.drawString(48,28,'6 October 2026  ·  v1.0  ·  Primary-source research; proposed design')
    c.drawRightString(w-48,28,str(d.page))

def paper():
    lines=(ROOT/'Enterprise_Agentic_Delivery_White_Paper.md').read_text().splitlines()
    story=[]
    story += [Spacer(1,70),Paragraph('Enterprise<br/>Agentic Delivery',ParagraphStyle(name='Cover',fontName='Body-Bold',fontSize=37,leading=45,textColor=colors.HexColor(NAVY))),Spacer(1,22),Paragraph('A small, enforceable architecture that grows through demonstrated outcomes',ParagraphStyle(name='Deck',fontName='Body',fontSize=17,leading=24,textColor=colors.HexColor(TEAL))),Spacer(1,36)]
    for t in ['A new centralized control-plane product is optional. Explicit authority, evidence, durable records and recovery remain necessary where consequences require them.','First-principles analysis · 13 architecture families · C1–C4 views · Value-gated construction','6 October 2026 · Version 1.0','This is a research-grounded architecture proposal. Product capabilities are documented; enterprise fitness and performance require implementation-specific validation.']:
        story+=[Paragraph(inline(t),styles['Text']),Spacer(1,9)]
    story += [PageBreak(),Paragraph('Contents',styles['Chapter'])]
    toc=TableOfContents(); toc.levelStyles=[styles['TOCEntry']]; story+=[toc,PageBreak()]
    i=0; skip_intro=True
    while i<len(lines):
        l=lines[i].strip()
        if l=='## Executive argument': skip_intro=False
        if skip_intro: i+=1; continue
        if not l: i+=1; continue
        if l.startswith('## '):
            if len(story)>5 and l!='## Executive argument': story.append(PageBreak())
            story.append(Paragraph(inline(l[3:]),styles['Chapter'])); i+=1; continue
        if l.startswith('### '):
            if re.match(r'### 4\.[3-6] ',l): story.append(PageBreak())
            story.append(Paragraph(inline(l[4:]),styles['Subhead'])); i+=1; continue
        if l.startswith('!['):
            path=re.search(r'\((.*?)\)',l).group(1)
            from PIL import Image as PILImage
            im=PILImage.open(ROOT/path); iw,ih=im.size
            width=490; height=width*ih/iw
            items=[Image(str(ROOT/path),width=width,height=height)]
            i+=1
            while i<len(lines) and not lines[i].strip(): i+=1
            if i<len(lines) and lines[i].startswith('**Figure'):
                items += [Spacer(1,6),Paragraph(inline(lines[i]),styles['Text'])]
                i+=1
            story.append(KeepTogether(items)); continue
        if l.startswith('```'):
            language=l[3:]; i+=1; block=[]
            while i<len(lines) and not lines[i].startswith('```'): block.append(lines[i]); i+=1
            i+=1
            if language=='mermaid': continue
            for s in block:
                story.append(Paragraph(escape(s).replace(' ','&#160;'),styles['ContractCode']))
            story.append(Spacer(1,10)); continue
        if l.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.match(r'^:?-+:?$',x) for x in cells): rows.append(cells)
                i+=1
            count=len(rows[0]); widths=[490/count]*count
            if count==2: widths=[170,320]
            if count==3: widths=[112,196,182]
            if count==4: widths=[85,140,130,135]
            if rows[0][0]=='Ref.': widths=[27,270,193]
            data=[[Paragraph(inline(x),styles['Cell']) for x in row] for row in rows]
            table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
            table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DCEAF2')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('TOPPADDING',(0,0),(-1,-1),7),('BOTTOMPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,0),.8,colors.HexColor(BLUE)),('LINEBELOW',(0,1),(-1,-1),.3,colors.HexColor('#D8E2E9')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F7F9FB')])]))
            story.extend([table,Spacer(1,10)]); continue
        paragraph=[l]; i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','```','![')):
            paragraph.append(lines[i].strip()); i+=1
        story.append(Paragraph(inline(' '.join(paragraph)),styles['Text']))
    d=PaperDoc(str(ROOT/'Enterprise_Agentic_Delivery_White_Paper.pdf'),pagesize=A4,rightMargin=52,leftMargin=52,topMargin=55,bottomMargin=48,title='Enterprise Agentic Delivery',author='Architecture research and synthesis')
    d.multiBuild(story,onFirstPage=page,onLaterPages=page)

def rgb(h): return RGBColor.from_string(h.lstrip('#'))
def textbox(slide,x,y,w,h,text,size=20,color=NAVY,bold=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True; tf.margin_left=0; tf.margin_right=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.size=Pt(size); p.font.name='Aptos'; p.font.color.rgb=rgb(color); p.font.bold=bold; p.space_after=Pt(9)
    return box

def slide_diagram(slide,name):
    nodes={}; edges=[]
    if name=='C1_Context.png':
        nodes={
          'owner':(4.6,2.1,4.1,.7,'Engineer / service owner / operator'),
          'system':(4.6,3.75,4.1,.9,'Enterprise Agentic Delivery'),
          'harness':(.55,5.7,3.8,.85,'Approved agent harnesses'),
          'identity':(9.0,5.7,3.8,.85,'Enterprise identity and policy'),
          'delivery':(4.6,5.7,4.1,.85,'Repository, build, release\nand operations systems')}
        edges=[('owner','system'),('system','harness'),('system','identity'),('system','delivery')]
    elif name=='C2_Containers.png':
        nodes={
          'client':(4.55,2.0,4.1,.65,'Existing IDE / portal / events'),
          'api':(4.55,3.15,4.1,.75,'Outcome Service / API'),
          'worker':(.55,3.15,3.45,.75,'Worker Gateway'),
          'runtime':(.55,4.6,3.45,.85,'Approved harness\nIsolated workspace'),
          'record':(9.1,2.0,3.6,.65,'Outcome database'),
          'evidence':(9.1,3.15,3.6,.75,'Evidence object store'),
          'action':(4.55,4.6,4.1,.85,'Trusted Action Gateway'),
          'policy':(9.1,4.6,3.6,.85,'IAM / policy authority'),
          'target':(4.55,6.0,4.1,.65,'Repository / release executors'),
          'durable':(9.1,6.0,3.6,.65,'Optional Workflow Runner')}
        edges=[('client','api'),('api','worker'),('worker','runtime'),('runtime','api'),('api','record'),('api','evidence'),('api','action'),('action','policy'),('action','target'),('target','api'),('api','durable')]
    elif name=='C3_Components.png':
        nodes={
          'intake':(4.55,2.0,4.1,.65,'Intake and Contract Validator'),
          'transition':(4.55,3.35,4.1,.85,'Transition Engine'),
          'planner':(.55,3.35,3.45,.85,'Planning and\nHarness Adapter'),
          'verifier':(.55,5.0,3.45,.85,'Evidence Validator'),
          'authority':(9.1,3.35,3.6,.85,'Policy and Approval\nConnector'),
          'effects':(9.1,5.0,3.6,.85,'Action Dispatcher\nand Reconciler'),
          'records':(4.55,5.0,4.1,.85,'Record Repository\nand Outbox'),
          'view':(4.55,6.3,4.1,.55,'Status Projection and Telemetry')}
        edges=[('intake','transition'),('transition','planner'),('planner','verifier'),('verifier','transition'),('transition','authority'),('authority','effects'),('effects','transition'),('transition','records'),('records','view')]
    else:
        nodes={
          'request':(4.55,2.0,4.1,.85,'ActionRequest\noperation / artifact / version / decision'),
          'dispatch':(4.55,3.8,4.1,.85,'ActionDispatcher\nsubmit / reconcile'),
          'policy':(.55,5.7,3.8,.85,'PolicyAuthority\nauthorize'),
          'executor':(4.6,5.7,4.1,.85,'ActionExecutor\napply / inspect'),
          'repo':(9.0,5.7,3.8,.85,'EffectRepository\nreserve / record')}
        edges=[('dispatch','request'),('dispatch','policy'),('dispatch','executor'),('dispatch','repo')]
    for a,b in edges:
        x1,y1,w1,h1,_=nodes[a]; x2,y2,w2,h2,_=nodes[b]
        c1=(x1+w1/2,y1+h1/2);c2=(x2+w2/2,y2+h2/2)
        dx,dy=c2[0]-c1[0],c2[1]-c1[1]
        t1=min(w1/2/abs(dx) if dx else 1e9,h1/2/abs(dy) if dy else 1e9)
        t2=min(w2/2/abs(dx) if dx else 1e9,h2/2/abs(dy) if dy else 1e9)
        p1=(c1[0]+dx*t1,c1[1]+dy*t1);p2=(c2[0]-dx*t2,c2[1]-dy*t2)
        if name=='C2_Containers.png' and (a,b) in [('api','durable'),('target','api')]:
            if b=='durable':
                points=[(8.65,3.53),(8.93,3.53),(8.93,5.72),(10.9,5.72),(10.9,6.0)]
            else:
                points=[(8.65,6.325),(8.78,6.325),(8.78,3.72),(8.65,3.72)]
            for j in range(len(points)-1):
                p,q=points[j:j+2]
                line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(p[0]),Inches(p[1]),Inches(q[0]),Inches(q[1]))
                line.line.color.rgb=rgb('#8096A8');line.line.width=Pt(1.1)
                if b=='durable':line.line.dash_style=MSO_LINE_DASH_STYLE.DASH
                if j==len(points)-2:
                    arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');line._element.spPr.get_or_add_ln().append(arrow)
            continue
        line=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(p1[0]),Inches(p1[1]),Inches(p2[0]),Inches(p2[1]))
        line.line.color.rgb=rgb('#8096A8');line.line.width=Pt(1.25)
        arrow=OxmlElement('a:tailEnd');arrow.set('type','triangle');line._element.spPr.get_or_add_ln().append(arrow)
    for key,(x,y,w,h,label) in nodes.items():
        sh=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        sh.fill.solid();core=key in ('system','api','action','transition','dispatch');sh.fill.fore_color.rgb=rgb(TEAL if core else LIGHT)
        sh.line.color.rgb=rgb(TEAL if core else '#8EAABD')
        tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=Inches(.09);tf.margin_right=Inches(.09);tf.margin_top=Inches(.06);tf.margin_bottom=0
        for i,t in enumerate(label.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=t;p.alignment=PP_ALIGN.CENTER;p.font.name='Aptos';p.font.size=Pt(16 if i==0 else 13);p.font.color.rgb=rgb('#FFFFFF' if core else NAVY)

def deck():
    prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
    blank=prs.slide_layouts[6]
    slides=[]
    def add(title,kicker,body=None,notes='',sources='',image=None,cols=None):
        s=prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('#FFFFFF')
        sh=s.shapes.add_shape(1,0,0,prs.slide_width,Inches(.12)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(TEAL); sh.line.fill.background()
        textbox(s,.6,.32,12,.4,kicker.upper(),12,TEAL,True)
        textbox(s,.6,.93,12,1.13,title,30,NAVY,True)
        if image:
            slide_diagram(s,image)
            notes += '\nPresenter emphasis:\n'+(body or '')
        elif cols:
            for i,(head,content) in enumerate(cols):
                x=.65+i*(12.0/len(cols)); w=11.2/len(cols)
                textbox(s,x,2.2,w,.58,head,20,TEAL,True)
                textbox(s,x,2.98,w,3.58,content,18)
        elif body: textbox(s,.7,2.3,11.9,4.1,body,23)
        textbox(s,.6,7.02,11.7,.2,sources or 'Architecture proposal · 6 October 2026 · See white paper for evidence and limits',9,GRAY)
        textbox(s,12.1,6.95,.55,.33,str(len(prs.slides)),11,GRAY)
        s.notes_slide.notes_text_frame.text=notes
        slides.append({'number':len(prs.slides),'title':title,'kicker':kicker,'notes':notes,'sources':sources})
        return s

    add('Enterprise Agentic Delivery','A small architecture that grows through outcomes',
        'I start with one useful outcome.\nI keep authority, evidence and recovery explicit.\nI add coordination only when a measured gap requires it.',
        notes='This deck follows the accompanying first-person white paper. It presents a design proposal, not a validated production platform. Product evidence was reviewed on 6 October 2026. The narrative begins with the challenge to control-plane necessity and ends with an evidence-gated investment decision.')
    add('A control-plane product is optional. Its responsibilities may still be needed.','The core finding',
        cols=[('Remove the mandatory platform','Existing IAM, repository rules, builds, release controllers and telemetry can supply control functions.'),('Keep consequences explicit','For consequential work, define authority, exact-change evidence, durable progress and recovery.')],
        notes='White paper §2. A counterexample disproves the universal claim that every enterprise workflow needs a new centralized control plane. A protected dependency-upgrade PR can succeed through existing systems. Conversely, an ambiguous production effect still requires authorization and reconciliation somewhere. This is a distinction between function and topology.')
    add('I tested both sides of the argument.','Disconfirmation',
        cols=[('A mandatory central plane','Disconfirmed by workflows already governed and recoverable through existing systems.'),('Harnesses remove all controls','Disconfirmed by stale approvals, uncertain external effects and conflicting workers.'),('A thin layer always helps','Reject it if it duplicates reliable existing state and costs more than it saves.')],
        notes='White paper §§2.2–2.4. My conclusion is conditional. I also challenge my own thin-layer recommendation: removal experiments must be permitted. Coordination is an investment, not an architectural virtue.')
    add('The harnesses offer complementary patterns.','Landscape / products',
        cols=[('Fleet and specialists','Cursor: direction separate from execution.\nCopilot: explicit dependencies.\nClaude: coding loop and bounded specialists.'),('Managed and missions','Codex: managed harness with environment choice.\nFactory: planned milestones and validators.\nKiro / AWS: development-to-operations access.')],
        sources='White paper references [1–11, 27–29] · Capabilities are documented; enterprise fit requires validation.',
        notes='White paper §§3.1–3.3. Cursor live documentation supports Enterprise with version 3.21.9 or later, unlike an older indexed result. Copilot Fleet SDK binding is experimental. Claude teams are confirmed, but detailed team semantics were not independently retrieved. Do not rank these products as universally strongest or infer production maturity from advertised scale.')
    add('Frameworks are alternatives, not layers to accumulate.','Landscape / frameworks and open designs',
        cols=[('State and composition','ADK: agents with deterministic nodes.\nLangGraph / Deep Agents: checkpoints and harness.\nMicrosoft Agent Framework: maintained successor to AutoGen.'),('Open and configurable','CrewAI: crews within flows.\nOpenHands: modular software-agent runtime.\nMetaGPT / ChatDev: explicit roles and configurable collaboration.')],
        sources='White paper references [12–21] · AutoGen is in maintenance mode.',
        notes='White paper §3.4. ChatDev 2.0 differs from its legacy virtual software company. Role simulation and workflow configuration are useful patterns, but are not evidence of enterprise production guarantees. Select one state owner and one primary harness initially.')
    add('Three contracts keep the architecture small.','The stable enterprise boundary',
        cols=[('Outcome','What I want.\nOwner, scope, constraints, acceptance, risk and budget.'),('Evidence','What was established.\nExact revision/digest, trusted producer, environment and freshness.'),('Action','What may change.\nBounded effect, authority, preconditions, operation identity and recovery.')],
        notes='White paper §4.1 and Appendix A. These contracts express enterprise semantics independently of a vendor’s hidden plan or session. A planning result is a proposal; an evidence result is a supported observation; an action authorization is permission for a bound external effect.')
    add('C1: the system fits into existing enterprise boundaries.','Architecture / context',image='C1_Context.png',
        body='People own outcomes.\nHarnesses reason and create.\nEnterprise systems authorize and execute.\nService observations establish results.',
        notes='White paper §4.3 / Figure 1. This is a C4 system-context view. The system can be domain operated. The harness is external because its lifecycle and trust boundary differ from the delivery system. References [22] establishes the C4 notation levels.')
    add('C2: reuse trusted executors; keep agents outside production authority.','Architecture / containers',image='C2_Containers.png',
        body='Outcome Service: add only if state has gaps.\nWorker Gateway: native harness integration.\nAction Gateway: existing executor can supply it.\nWorkflow Runner: optional.',
        notes='White paper §§4.2 and 4.4 / Figure 2. Target logical implementation, not the minimum Stage 1 deployment. A gateway is ineffective if agents retain alternate privileged routes. Credentials, management-network paths and protected verification jobs must enforce the boundary. Database/object store roles should reuse established services.')
    add('C3: deterministic transitions surround flexible planning.','Architecture / components',image='C3_Components.png',
        body='Planning proposes.\nVerification checks exact evidence.\nPolicy authorizes.\nReconciliation resolves uncertain effects.\nStart as modules in one application.',
        notes='White paper §4.5 / Figure 3. This zooms into the Outcome Service container. Separate modules do not imply microservices. Enterprise readiness depends on ownership and enforceable behavior rather than the number of boxes.')
    add('C4: the code boundary handles uncertainty explicitly.','Architecture / code',image='C4_Code.png',
        body='Reserve a unique operation.\nRecheck current authority.\nApply with executor guarantees.\nInspect after lost response.\nDo not retry unknown non-idempotent effects.',
        notes='White paper §4.6 / Figure 4. Planned interfaces, not deployed code. ActionDispatcher depends on ActionRequest, PolicyAuthority, ActionExecutor and EffectRepository. Use request hashes, optimistic versions, outbox/inbox, operation IDs and fencing where supported. There is no universal exactly-once transaction across arbitrary tools.')
    add('Completion means the intended outcome was verified.','Lifecycle',
        body='A paper completes when research and review criteria pass.\nA PR completes when accepted behavior is demonstrated.\nA release completes after its health contract passes.\nAn uncertain effect remains Indeterminate until reconciled.',
        notes='White paper §§5.1, 7.3. Agent completion and business completion are separate. Defined, Active, Waiting, Ready for Action, Acting, Observing, Completed, Failed, Cancelled and Indeterminate are coarse outcome states. Domain acceptance defines their criteria.')
    add('Run useful checks early; retain independent enforcement.','Evidence',
        cols=[('At creation','Compile, test contracts, scan dependencies and apply policy before costly shared execution.'),('At acceptance','Bind trusted evidence to revision, artifact, suite, environment and policy. Reject stale or mismatched proof.')],
        notes='White paper §§6.1–6.2. Local checks reduce rework but do not automatically become release evidence. Protected suites and independent producers prevent the builder from defining success through its own assertions. SLSA provenance provides build lineage, not a correctness guarantee. Source [25].')
    add('The hard cases are external effects and bypass paths.','Failure and security',
        cols=[('Lost response','Inspect the executor operation before retrying. If the result is unknown, stop and reconcile.'),('Stale authority','Bind approvals to digest, target, operation, policy and expiry. Fence stale workers.'),('Prompt failure','Deny privileged paths outside the trusted executor. Treat retrieved instructions as untrusted content.')],
        notes='White paper §§6.3 and 7.1. Also test duplicate events, unavailable policy, unavailable telemetry, cancellation with in-flight effects, and cross-domain access. A controller may finish a previously authorized bounded safety sequence during an outage; new privileged actions remain blocked.')
    add('Memory improves work only if its learning can be evaluated and reversed.','Governed learning',
        body='Keep authoritative workflow state out of semantic memory.\nPromote verified procedures with source, scope, owner and expiry.\nEvaluate instruction changes on held-out cases.\nUse a small cohort and retain the prior version.',
        notes='White paper §6.4. Vector retrieval is optional. Access checks, freshness, deletion and demonstrable improvement justify adding it. A successful run can propose knowledge; it does not grant permission to rewrite policy.')
    add('Stages 0–2 already deliver useful results.','Build through value gates',
        cols=[('0 / Understand','Baseline one frequent outcome.\nIdentify tests, authority and actual waiting costs.\nUse a script if sufficient.'),('1 / Improve creation','One harness, one delivery path.\nBetter accepted PRs and lower human effort.\nExisting controls remain authoritative.'),('2 / Recover work','Add missing shared records and reconciliation.\nReduce interruption and coordination effort.')],
        notes='White paper §§8.1–8.3. Illustrative pilot: 20–30 comparable tasks where diversity permits, and a proposed 20% median human-effort reduction without material quality decline. These are proposed targets, not empirical forecasts. Small cohorts cannot prove rare-event safety.')
    add('Stages 3–5 expand only after evidence supports them.','Build through value gates',
        cols=[('3 / Scale bounded work','Compare single-agent against owned parallel tasks.\nMigrate in repository waves.\nStop if coupling erases benefits.'),('4 / Validate releases','One trusted executor and health contract.\nShadow first, bounded cohort next.\nTest pause and recovery.'),('5 / Broaden outcomes','Investigate incidents before writes.\nAdd domain packs and governed learning.\nEvaluate every new outcome family.')],
        notes='White paper §§8.4–8.7. No fixed delivery dates without team and readiness information. Every stage has standalone value and stopping criteria. Central portfolio scheduling appears only if observed domain coordination requires it.')
    add('The same contracts support many enterprise outcomes.','Mature scope',
        cols=[('Engineering','Research and architecture\nFeatures and defect repair\nDependencies and migrations\nPolicy and pipeline remediation'),('Delivery and operations','Validated release\nIncident investigation\nBounded recovery\nInfrastructure and data changes')],
        notes='White paper §7.3. Each family has distinct evidence, risk, and recovery rules. Data and irreversible migration actions require additional domain validation. Success with upgrades does not authorize production database repair.')
    add('CI/CD products retire only after their functions are replaced safely.','Migration',
        body='Start with Jenkins, Spinnaker, GitLab or other existing executors.\nRequest contract-shaped outcomes through adapters.\nReplace a tool only after equivalent controls, recovery and economics pass.\nAn agent loop does not remove the need for reproducible execution.',
        notes='White paper §10. Distinguish product redundancy from function redundancy. Evaluate isolation, provenance, concurrency, cancellation, authorization, observation and recovery before retiring an incumbent. Keep one reversible workflow migration route.')
    add('Measure accepted outcomes, total effort and service effects.','Economics',
        cols=[('Useful measures','Accepted outcomes / all attempts\nValidated completion time\nHuman effort and review rejection\nEscaped defects and service regressions'),('Full cost','Model and compute\nIntegration and rework\nHuman review\nAllocated platform operation\nDivided by accepted outcomes')],
        notes='White paper §9. PR count alone is not value. Include failed and abandoned tasks. Distinguish commit-to-production lead time from production workflow duration. The paper’s 80-hour net savings example is illustrative arithmetic, not a forecast.')
    add('I would fund one pilot and preserve the right to simplify.','Decision',
        body='One frequent outcome. One approved harness. One trusted path.\nDefine evidence and authority before expanding the fleet.\nKeep the custom layer only where it proves useful.\nProgress by measured value and tested recovery.',
        notes='White paper §§11–13. Test removal of the custom service, single-agent versus fleet, native state versus outer workflow, memory versus no memory, and incumbent versus replacement execution. These tests can overturn the design. The recommendation is an actionable pilot, not a request to build every target component.')
    add('Evidence limits are part of the recommendation.','Research boundary',
        body='Primary documentation reviewed on 6 October 2026.\nNo paid-product benchmark or production pilot was performed.\nProduct access and versions change; recheck before procurement.\nEnterprise obligations, RTO/RPO and ROI remain deployment-specific.',
        notes='White paper §1.3 and Appendix C. The slide deck and paper are a cohesive design proposal. The live Cursor page superseded an indexed restriction. AutoGen maintenance status was checked in the maintained repository. Claude team existence is verified; detailed team lifecycle guarantees remain a validation question.')
    source_rows=[]
    md=(ROOT/'Enterprise_Agentic_Delivery_White_Paper.md').read_text()
    source_text=md.split('## Appendix C. Sources and claim provenance',1)[1]
    for match in re.finditer(r'^\| (\d+) \| (.*?) \| (.*?) \|$',source_text,re.M):
        source_rows.append((match[1],match[2],match[3]))
    for start in [0,15]:
        subset=source_rows[start:start+15]
        s=add('Primary-source ledger','References '+str(start+1)+'–'+str(start+len(subset)),notes='Full URLs, supported claims and implementation questions are in white paper Appendix C. All references were consulted on 6 October 2026.\n'+'\n'.join(f'[{n}] {src} — {claim}' for n,src,claim in subset))
        for col in range(2):
            rows=subset[col*8:(col+1)*8]
            text='\n'.join(f'[{n}] '+src.split(' — ')[0] for n,src,claim in rows)
            textbox(s,.7+col*6.3,2.1,5.8,4.5,text,17)
    prs.save(ROOT/'Enterprise_Agentic_Delivery_Slides.pptx')
    (ROOT/'slide_manifest.json').write_text(json.dumps(slides,indent=2))

if '--diagrams-only' not in sys.argv:
    paper(); deck()
    print('Created paper PDF, PPTX, 4 PNGs and 4 editable SVGs.')
else:
    print('Created 4 validated PNGs and 4 editable SVGs.')
